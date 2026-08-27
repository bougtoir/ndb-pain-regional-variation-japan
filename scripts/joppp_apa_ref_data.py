#!/usr/bin/env python3
"""APA 7th (T&F standard) reference data and citation helpers for the JoPPP manuscript."""

import re
from docx.shared import Pt, Inches

REF_DATA = [
    {
        'num': 1,
        'surnames': ['Callister'],
        'year': 2003,
        'apa': """Callister, L. C. (2003). Cultural influences on pain perceptions and behaviors. *Home Health Care Management & Practice*, *15*(3), 207\u2013211. https://doi.org/10.1177/1084822302250687""",
    },
    {
        'num': 2,
        'surnames': ['Rogger', 'Bello', 'Romero', 'Urman', 'Luedi', 'Filipovic'],
        'year': 2023,
        'apa': """Rogger, R., Bello, C., Romero, C. S., Urman, R. D., Luedi, M. M., & Filipovic, M. G. (2023). Cultural framing and the impact on acute pain and pain services. *Current Pain and Headache Reports*, *27*(9), 429\u2013436. https://doi.org/10.1007/s11916-023-01125-2""",
    },
    {
        'num': 3,
        'surnames': ['Zborowski'],
        'year': 1969,
        'apa': """Zborowski, M. (1969). *People in pain*. Jossey-Bass.""",
    },
    {
        'num': 4,
        'surnames': ['Okolo', 'Olorunsogo', 'Babawarun'],
        'year': 2024,
        'apa': """Okolo, C. A., Olorunsogo, T., & Babawarun, O. (2024). Cultural variability in pain perception: A review of cross-cultural studies. *International Journal of Science and Research Archive*, *11*(1), 2550\u20132556. https://doi.org/10.30574/ijsra.2024.11.1.0339""",
    },
    {
        'num': 5,
        'surnames': ['Hobara'],
        'year': 2005,
        'apa': """Hobara, M. (2005). Beliefs about appropriate pain behavior: Cross-cultural and sex differences between Japanese and Euro-Americans. *European Journal of Pain*, *9*(4), 389\u2013393. https://doi.org/10.1016/j.ejpain.2004.09.006""",
    },
    {
        'num': 6,
        'surnames': ['Feng', 'Herdman', 'van Nooten', 'Cleeland', 'Parkin', 'Ikeda', 'Igarashi', 'Devlin'],
        'year': 2017,
        'apa': """Feng, Y., Herdman, M., van Nooten, F., Cleeland, C., Parkin, D., Ikeda, S., Igarashi, A., & Devlin, N. J. (2017). An exploration of differences between Japan and two European countries in the self-reporting and valuation of pain and discomfort on the EQ-5D. *Quality of Life Research*, *26*(8), 2067\u20132078. https://doi.org/10.1007/s11136-017-1541-5""",
    },
    {
        'num': 7,
        'surnames': ['Cohen', 'Nisbett', 'Bowdle', 'Schwarz'],
        'year': 1996,
        'apa': """Cohen, D., Nisbett, R. E., Bowdle, B. F., & Schwarz, N. (1996). Insult, aggression, and the southern culture of honor: An "experimental ethnography." *Journal of Personality and Social Psychology*, *70*(5), 945\u2013960. https://doi.org/10.1037/0022-3514.70.5.945""",
    },
    {
        'num': 8,
        'surnames': ['Kumagai'],
        'year': 2020,
        'apa': """Kumagai, S. (2020). \u30e1\u30c7\u30a3\u30a2\u304c\u518d\u751f\u7523\u3059\u308b\u6771\u5317\u50cf\uff1a\u300e\u79d8\u5bc6\u306e\u30b1\u30f3\u30df\u30f3SHOW\u300f\u3067\u306e\u6771\u5317\u5fa9\u8208\u30b3\u30fc\u30ca\u30fc\u3092\u3081\u3050\u3063\u3066 [Media representations reproducing images of Tohoku: The Tohoku reconstruction corner in "Secret Kenmin SHOW"]. *\u3053\u3068\u3070* (Kotoba), *41*, 21\u201338. https://doi.org/10.20741/kotoba.41.0_21""",
    },
    {
        'num': 9,
        'surnames': ['Takeda', 'Yarimizu'],
        'year': 2016,
        'apa': """Takeda, K., & Yarimizu, K. (2016). \u75db\u307f\u3092\u8868\u3059\u8a00\u8a9e\u8868\u73fe\u30a6\u30ba\u30af\u306e\u5730\u57df\u5dee [Regional differences in the pain expression "uzuku"]. *NINJAL Research Papers*, *10*, 221\u2013243. https://doi.org/10.15084/00000816""",
    },
    {
        'num': 10,
        'surnames': [],
        'short': 'Pfizer Japan Inc.',
        'year': 2017,
        'apa': """Pfizer Japan Inc. (2017). 47-prefecture survey on chronic pain (2012 vs 2017 comparison) [in Japanese]. Research Research. https://www.lisalisa50.com/research20171014_2.html""",
    },
    {
        'num': 11,
        'surnames': [],
        'short': 'MHLW',
        'year': None,
        'apa': """Ministry of Health, Labour and Welfare. (n.d.). NDB Open Data, 10th edition. Retrieved January 15, 2025, from https://www.mhlw.go.jp/stf/seisakunitsuite/bunya/0000177221_00016.html""",
    },
    {
        'num': 12,
        'surnames': ['Wakaizumi', 'Tanaka', 'Shinohara', 'Wu', 'Takaoka', 'Kawate', 'Oka', 'Matsudaira'],
        'year': 2024,
        'apa': """Wakaizumi, K., Tanaka, C., Shinohara, Y., Wu, Y., Takaoka, S., Kawate, M., Oka, H., & Matsudaira, K. (2024). Geographical variation in high-impact chronic pain and psychological associations at the regional level: A multilevel analysis of a large-scale internet-based cross-sectional survey. *Frontiers in Public Health*, *12*, Article 1482177. https://doi.org/10.3389/fpubh.2024.1482177""",
    },
    {
        'num': 13,
        'surnames': ['Takahashi', 'Miyashita', 'Nakazawa', 'Wada', 'Matsuoka'],
        'year': 2025,
        'apa': """Takahashi, R., Miyashita, M., Nakazawa, Y., Wada, S., & Matsuoka, Y. (2025). Population-based claims study of regional and hospital function differences in opioid prescribing for cancer patients who died in hospital in Japan. *Japanese Journal of Clinical Oncology*, *55*(12), 1372\u20131377. https://doi.org/10.1093/jjco/hyaf149""",
    },
    {
        'num': 14,
        'surnames': ['Taira', 'Mori', 'Ishimaru', 'Iwagami', 'Sakata', 'Watanabe', 'Takahashi', 'Tamiya'],
        'year': 2021,
        'apa': """Taira, K., Mori, T., Ishimaru, M., Iwagami, M., Sakata, N., Watanabe, T., Takahashi, H., & Tamiya, N. (2021). Regional inequality in dental care utilization in Japan: An ecological study using the National Database of Health Insurance Claims. *The Lancet Regional Health - Western Pacific*, *12*, Article 100170. https://doi.org/10.1016/j.lanwpc.2021.100170""",
    },
    {
        'num': 15,
        'surnames': ['von Elm', 'Altman', 'Egger', 'Pocock', 'G\u00f8tzsche', 'Vandenbroucke'],
        'year': 2007,
        'apa': """von Elm, E., Altman, D. G., Egger, M., Pocock, S. J., G\u00f8tzsche, P. C., & Vandenbroucke, J. P. (2007). The Strengthening the Reporting of Observational Studies in Epidemiology (STROBE) statement: Guidelines for reporting observational studies. *The Lancet*, *370*(9596), 1453\u20131457. https://doi.org/10.1016/S0140-6736(07)61602-X""",
    },
    {
        'num': 16,
        'surnames': ['Benchimol', 'Smeeth', 'Guttmann', 'Harron', 'Moher', 'Petersen', 'S\u00f8rensen', 'von Elm', 'Langan'],
        'year': 2015,
        'apa': """Benchimol, E. I., Smeeth, L., Guttmann, A., Harron, K., Moher, D., Petersen, I., S\u00f8rensen, H. T., von Elm, E., Langan, S. M., & RECORD Working Committee. (2015). The REporting of studies Conducted using Observational Routinely-collected health Data (RECORD) Statement. *PLoS Medicine*, *12*(10), Article e1001885. https://doi.org/10.1371/journal.pmed.1001885""",
    },
    {
        'num': 17,
        'surnames': ['Anderson', 'Green', 'Payne'],
        'year': 2009,
        'apa': """Anderson, K. O., Green, C. R., & Payne, R. (2009). Racial and ethnic disparities in pain: Causes and consequences of unequal care. *The Journal of Pain*, *10*(12), 1187\u20131204. https://doi.org/10.1016/j.jpain.2009.10.002""",
    },
    {
        'num': 18,
        'surnames': ['Campbell', 'Edwards'],
        'year': 2012,
        'apa': """Campbell, C. M., & Edwards, R. R. (2012). Ethnic differences in pain and pain management. *Pain Management*, *2*(3), 219\u2013230. https://doi.org/10.2217/pmt.12.7""",
    },
    {
        'num': 19,
        'surnames': ['Befu'],
        'year': 2001,
        'apa': """Befu, H. (2001). *Hegemony of homogeneity: An anthropological analysis of Nihonjinron*. Trans Pacific Press.""",
    },
    {
        'num': 20,
        'surnames': ['Burgess'],
        'year': 2010,
        'apa': """Burgess, C. (2010). The "illusion" of homogeneous Japan and national character: Discourse as a tool to transcend the "myth" vs. "reality" binary. *Asia-Pacific Journal*, *8*(9), 1\u201322. https://doi.org/10.1017/s1557466010009381""",
    },
    {
        'num': 21,
        'surnames': ['Raja', 'Carr', 'Cohen', 'Finnerup', 'Flor', 'Gibson', 'Keefe', 'Mogil', 'Ringkamp', 'Sluka'],
        'year': 2020,
        'apa': """Raja, S. N., Carr, D. B., Cohen, M., Finnerup, N. B., Flor, H., Gibson, S., Keefe, F. J., Mogil, J. S., Ringkamp, M., & Sluka, K. A. (2020). The revised International Association for the Study of Pain definition of pain: Concepts, challenges, and compromises. *Pain*, *161*(9), 1976\u20131982. https://doi.org/10.1097/j.pain.0000000000001939""",
    },
    {
        'num': 22,
        'surnames': ['Kehlet', 'Jensen', 'Woolf'],
        'year': 2006,
        'apa': """Kehlet, H., Jensen, T. S., & Woolf, C. J. (2006). Persistent postsurgical pain: Risk factors and prevention. *The Lancet*, *367*(9522), 1618\u20131625. https://doi.org/10.1016/S0140-6736(06)68700-X""",
    },
]

for _r in REF_DATA:
    if _r.get('short') and not _r['surnames']:
        _r['first_author'] = _r['short']
    else:
        _r['first_author'] = _r['surnames'][0]

ref_data = REF_DATA


def _apa_in_text(ref):
    surnames = ref.get('surnames', [])
    n = len(surnames)
    if n == 0:
        return ref.get('short', ref.get('first_author', ''))
    if n == 1:
        return surnames[0]
    if n == 2:
        return f'{surnames[0]} & {surnames[1]}'
    return f'{surnames[0]} et al.'


def cite(*nums):
    """Return an APA parenthetical citation string for the given reference numbers."""
    entries = [ref_data[n - 1] for n in nums]
    entries.sort(key=lambda r: (
        _apa_in_text(r).lower().replace(' & ', ' '),
        r.get('year') or 0,
    ))
    parts = []
    for e in entries:
        year = str(e['year']) if e.get('year') else 'n.d.'
        parts.append(f'{_apa_in_text(e)}, {year}')
    return '(' + '; '.join(parts) + ')'


def add_ref_runs(p, text):
    """Parse text with {n} or {n,m} markers and create runs with font-based superscript."""
    parts = re.split(r'(\{[^}]+\})', text)
    for part in parts:
        if part.startswith('{') and part.endswith('}'):
            run = p.add_run(part[1:-1])
            run.font.superscript = True
            run.font.size = Pt(10)
        else:
            p.add_run(part)


def add_apa_ref_paragraph(p, text):
    """Add an APA reference-list entry, italicising text between asterisks."""
    parts = re.split(r'(\*[^*]+\*)', text)
    for part in parts:
        if part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.italic = True
        else:
            p.add_run(part)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)